import csv
import io
import json
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django import forms
from django.contrib import admin, messages
from django.core.exceptions import PermissionDenied
from django.http import HttpResponseRedirect
from django.shortcuts import render
from django.urls import path, reverse
from django.utils.html import format_html
from openpyxl import load_workbook
from .models import (
    BirthdayWish,
    CompanyUnit,
    Department,
    DepartmentScheduleChange,
    Employee,
    EmployeeAssignmentHistory, EmployeeStatusHistory,
    EmployeeDependent,
    Position,
    SocialInsurancePeriod,
    SOCIAL_INSURANCE_PARTICIPATION_STATUS,
)


class EmployeeSalaryCSVUploadForm(forms.Form):
    file = forms.FileField(label='File CSV/TSV')


class EmployeeLegalEntityCSVUploadForm(forms.Form):
    file = forms.FileField(label='File CSV/TSV')


class EmployeeSocialInsuranceUploadForm(forms.Form):
    file = forms.FileField(label='File Excel (.xlsx) hoặc CSV/TSV')


def _normalize_employee_code(raw_code):
    if raw_code is None:
        return ''
    normalized = str(raw_code).strip().replace('\ufeff', '').replace('\xa0', '').upper()
    if normalized.endswith(',00') or normalized.endswith('.00'):
        normalized = normalized[:-3]
    normalized = normalized.replace(' ', '')
    return normalized


def _parse_money(raw_value, default=None):
    if raw_value is None:
        return default

    text = str(raw_value).strip().replace('\xa0', '').replace(' ', '')
    if not text:
        return default

    if text in {'-', '--'}:
        return default

    if '.' in text and ',' in text:
        if text.rfind(',') > text.rfind('.'):
            text = text.replace('.', '').replace(',', '.')
        else:
            text = text.replace(',', '')
    elif ',' in text:
        parts = text.split(',')
        if len(parts) == 2 and len(parts[1]) <= 2:
            text = '.'.join(parts)
        else:
            text = ''.join(parts)
    elif '.' in text:
        parts = text.split('.')
        if len(parts) > 2 or (len(parts) == 2 and len(parts[1]) == 3 and parts[0].isdigit()):
            text = ''.join(parts)

    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        raise ValueError(f"Giá trị tiền tệ không hợp lệ: {raw_value}")


def _get_cell(row, *keys):
    normalized_map = {str(k).strip().lower(): v for k, v in row.items()}
    for key in keys:
        value = normalized_map.get(key.lower())
        if value is not None:
            return value
    return None


def _normalize_optional_text(raw_value):
    if raw_value is None:
        return None
    if isinstance(raw_value, str):
        value = raw_value.strip()
    else:
        value = str(raw_value).strip()
    return value or None


def _parse_date_value(raw_value, default=None):
    if raw_value is None:
        return default

    if isinstance(raw_value, datetime):
        return raw_value.date()
    if isinstance(raw_value, date):
        return raw_value

    text = str(raw_value).strip()
    if not text or text in {'-', '--'}:
        return default

    date_formats = (
        '%Y-%m-%d',
        '%d/%m/%Y',
        '%d-%m-%Y',
        '%m/%d/%Y',
        '%Y/%m/%d',
    )

    for fmt in date_formats:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue

    try:
        return date.fromisoformat(text)
    except ValueError as exc:
        raise ValueError(f"Ngày không hợp lệ: {raw_value}") from exc


def _read_tabular_rows(upload_file):
    filename = (upload_file.name or '').lower()
    if filename.endswith(('.xlsx', '.xlsm')):
        raw_bytes = upload_file.read()
        workbook = load_workbook(filename=io.BytesIO(raw_bytes), data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []

        headers = [str(cell).strip() if cell is not None else '' for cell in rows[0]]
        if not any(headers):
            raise ValueError('Không đọc được header của file Excel.')

        normalized_headers = [header if header else f'column_{idx}' for idx, header in enumerate(headers)]
        output = []
        for row in rows[1:]:
            row_values = list(row)
            if not any(value not in (None, '') for value in row_values):
                continue
            output.append({
                normalized_headers[idx]: row_values[idx] if idx < len(row_values) else None
                for idx in range(len(normalized_headers))
            })
        return output

    raw_bytes = upload_file.read()
    text = raw_bytes.decode('utf-8-sig', errors='ignore')
    sample = text[:2000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters='\t,;')
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = '\t' if '\t' in sample else ','

    reader = csv.DictReader(text.splitlines(), delimiter=delimiter)
    if not reader.fieldnames:
        raise ValueError('Không đọc được header của file CSV.')
    return list(reader)

# -----------------------------
# CompanyUnit Admin
# -----------------------------
@admin.register(CompanyUnit)
class CompanyUnitAdmin(admin.ModelAdmin):
    list_display = ('code', 'prefix_code', 'name', 'is_active', 'created_at')
    list_filter = ('is_active',)
    search_fields = ('code', 'prefix_code', 'name')
    ordering = ('name',)


# -----------------------------
# Department Admin
# -----------------------------
class DepartmentScheduleChangeInline(admin.TabularInline):
    """Đổi lịch làm việc của phòng ban từ một ngày — employee/schedule.py."""
    model = DepartmentScheduleChange
    extra = 0
    fields = ('effective_from', 'works_on_sunday', 'saturday_afternoon_always_workday',
              'online_weekend_unlimited', 'note')
    verbose_name = "Đổi lịch làm việc từ ngày"
    verbose_name_plural = (
        "Đổi lịch làm việc theo ngày hiệu lực — mỗi dòng ghi lịch MỚI (đủ 3 cờ) áp dụng từ "
        "ngày đó; các cờ ở mục \"Cấu hình chấm công\" là lịch GỐC cho các ngày trước lần đổi "
        "đầu tiên, KHÔNG sửa chúng khi đổi lịch"
    )


class DepartmentAdmin(admin.ModelAdmin):
    list_display = ('code', 'name', 'is_section', 'parent_department', 'manager_info', 'employee_count', 'works_on_sunday', 'saturday_afternoon_always_workday', 'is_deleted', 'created_at')
    list_filter = ('is_section', 'is_deleted', 'parent_department', 'works_on_sunday', 'saturday_afternoon_always_workday', 'online_weekend_unlimited')
    search_fields = ('code', 'name', 'description')
    raw_id_fields = ('parent_department', 'manager', 'deleted_by')
    ordering = ('code',)
    inlines = [DepartmentScheduleChangeInline]

    fieldsets = (
        ('Thông tin cơ bản', {
            'fields': ('code', 'name', 'description', 'parent_department', 'is_section')
        }),
        ('Quản lý', {
            'fields': ('manager',)
        }),
        ('Cấu hình chấm công', {
            'fields': ('works_on_sunday', 'saturday_afternoon_always_workday', 'online_weekend_unlimited'),
            'description': (
                '"Có đi làm Chủ nhật": mặc định TẮT (phòng ban nghỉ Chủ nhật, Chủ nhật không '
                'bị tính là "nghỉ không tạo đơn"). Chỉ bật cho phòng thực sự làm Chủ nhật. '
                'Mặc định thứ 7 tuần 2/4 trong tháng chỉ làm nửa ngày (chiều không tự '
                'tính công, phải làm đơn tăng ca). Bật nếu phòng ban làm chiều thứ 7 mọi tuần. '
                'Đây là lịch GỐC: muốn đổi lịch TỪ MỘT NGÀY (không áp ngược cho các ngày trước) '
                'thì thêm dòng ở bảng "Đổi lịch làm việc theo ngày hiệu lực" bên dưới.'
            ),
        }),
        ('Xoá (soft-delete)', {
            'fields': ('is_deleted', 'deleted_at', 'deleted_by'),
            'description': (
                'Xoá qua trang Quản lý phòng ban/Quản lý bộ phận chỉ đánh dấu is_deleted '
                '(không mất dữ liệu) — bỏ tick is_deleted ở đây để khôi phục lại bản ghi.'
            ),
        }),
    )
    
    def manager_info(self, obj):
        if obj.manager:
            return format_html(
                '{} ({})',
                obj.manager.full_name,
                obj.manager.employee_id
            )
        return "Chưa có"
    manager_info.short_description = "Trưởng phòng"
    
    def employee_count(self, obj):
        return obj.employees.count()
    employee_count.short_description = "Số nhân viên"


# -----------------------------
# Position Admin
# -----------------------------
class PositionAdmin(admin.ModelAdmin):
    list_display = ('code', 'title', 'departments_display', 'level', 'is_management', 'overtime_grace_deadline_enabled', 'parent_position', 'employee_count')
    list_filter = ('department', 'is_management', 'level', 'overtime_grace_deadline_enabled')
    search_fields = ('code', 'title', 'description', 'department__name', 'department__code')
    raw_id_fields = ('parent_position',)
    filter_horizontal = ('department',)
    ordering = ('level', 'code')

    fieldsets = (
        ('Thông tin cơ bản', {
            'fields': ('code', 'title', 'description', 'department')
        }),
        ('Cấp bậc', {
            'fields': ('level', 'parent_position', 'is_management')
        }),
        ('Gia hạn đơn tăng ca (bất khả kháng)', {
            'fields': ('overtime_grace_deadline_enabled',),
            'description': (
                'Bật cho vị trí có thể phát sinh tăng ca kéo dài quá giờ ngoài kế '
                'hoạch (vd ca mổ không thể dừng đúng giờ). Hạn TẠO và DUYỆT đơn tăng ca '
                'của vị trí này nới tới 10h00 sáng ngày làm việc kế tiếp, thay vì phải '
                'xong trước 2 giờ so với giờ bắt đầu tăng ca.'
            ),
        }),
    )
    
    def employee_count(self, obj):
        return Employee.objects.filter(position=obj).count()
    employee_count.short_description = "Số nhân viên"

    def departments_display(self, obj):
        return ", ".join(obj.department.values_list('name', flat=True))
    departments_display.short_description = "Phòng ban"


# -----------------------------
# Employee Admin Form  (exposes extra_info keys as real fields)
# -----------------------------
SUB_DEPARTMENT_CHOICES = [
    ('', '---------'),
    ('ADS', 'ADS'), ('ADS2', 'ADS2'),
    ('KD 1', 'KD 1'), ('KD 2', 'KD 2'), ('KD 3', 'KD 3'), ('KD 4', 'KD 4'),
    ('KD 5', 'KD 5'), ('KD MN', 'KD MN'),
    ('CSKH MB', 'CSKH MB'), ('CSKH MN', 'CSKH MN'),
    ('Media MB', 'Media MB'), ('Media MN', 'Media MN'), ('Media', 'Media'),
    ('Giám sát nội bộ', 'Giám sát nội bộ'),
    ('HCNS', 'HCNS'), ('Kế toán', 'Kế toán'), ('Truyền thông', 'Truyền thông'),
    ('TTTH 01', 'TTTH 01'), ('TTTH 02', 'TTTH 02'), ('TTTH 03', 'TTTH 03'),
    ('TTTH 04', 'TTTH 04'), ('TTTH 05', 'TTTH 05'), ('TTTH 06', 'TTTH 06'),
    ('TTTH 07', 'TTTH 07'), ('TTTH 08', 'TTTH 08'), ('TTTH 09', 'TTTH 09'),
    ('TTTH 10', 'TTTH 10'), ('TTTH 11', 'TTTH 11'), ('TTTH 15', 'TTTH 15'),
    ('Xây group', 'Xây group'),
    ('Tiktok 1', 'Tiktok 1'), ('Tiktok 2', 'Tiktok 2'), ('Tiktok 3', 'Tiktok 3'),
    ('Tiktok 4', 'Tiktok 4'), ('Tiktok 5', 'Tiktok 5'), ('Tiktok 6', 'Tiktok 6'),
    ('Tiktok 7', 'Tiktok 7'), ('Tiktok 8', 'Tiktok 8'), ('Tiktok 9', 'Tiktok 9'),
    ('Tiktok 10', 'Tiktok 10'),
    ('Pháp chế', 'Pháp chế'), ('Mua hàng', 'Mua hàng'),
    ('IT', 'IT'), ('AI', 'AI'), ('Công nghệ thông tin', 'Công nghệ thông tin'),
]


class EmployeeAdminForm(forms.ModelForm):
    """Custom form to expose extra_info JSON keys as proper admin fields."""

    @staticmethod
    def _normalize_extra_info(value):
        if isinstance(value, dict):
            return value
        if isinstance(value, str):
            raw = value.strip()
            if not raw:
                return {}
            try:
                parsed = json.loads(raw)
            except (TypeError, ValueError):
                return {}
            return parsed if isinstance(parsed, dict) else {}
        return {}

    work_type = forms.CharField(
        required=False,
        label='Loại công việc (work_type)',
        help_text='Lưu trong extra_info. VD: Hành chính, Kỹ thuật…',
        widget=forms.TextInput(attrs={'size': 40}),
    )
    sub_department = forms.ChoiceField(
        required=False,
        label='Bộ phận con (sub_department)',
        help_text='Lưu trong extra_info.',
        choices=SUB_DEPARTMENT_CHOICES,
    )
    team = forms.CharField(
        required=False,
        label='Team',
        help_text='Lưu trong extra_info.',
        widget=forms.TextInput(attrs={'size': 40}),
    )

    class Meta:
        model = Employee
        fields = '__all__'

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        if self.instance and self.instance.pk:
            extra = self._normalize_extra_info(self.instance.extra_info)
            self.fields['work_type'].initial = extra.get('work_type', '')
            self.fields['sub_department'].initial = extra.get('sub_department', '')
            self.fields['team'].initial = extra.get('team', '')

    def save(self, commit=True):
        instance = super().save(commit=False)
        extra = dict(self._normalize_extra_info(instance.extra_info))
        for key in ('work_type', 'sub_department', 'team'):
            val = self.cleaned_data.get(key, '')
            if val:
                extra[key] = val
            else:
                extra.pop(key, None)
        instance.extra_info = extra
        if commit:
            instance.save()
            self.save_m2m()
        return instance


# -----------------------------
# Employee Admin
# -----------------------------
class EmployeeAssignmentHistoryInline(admin.TabularInline):
    """Lịch sử điều chuyển — employee/assignment.py."""
    model = EmployeeAssignmentHistory
    extra = 0
    fields = ('changed_from', 'department', 'position', 'doctor_team', 'work_form', 'personal_shift_ids', 'note')
    verbose_name = "Dữ liệu cũ trước ngày"
    verbose_name_plural = (
        "Lịch sử điều chuyển — tự ghi khi đổi phòng ban/vị trí/Team bác sĩ/hình thức làm việc/ca "
        "cá nhân. Mỗi dòng là dữ liệu CŨ, dùng tính công cho các ngày TRƯỚC \"Dữ liệu mới áp dụng từ "
        "ngày\" (mặc định ngày lưu; sửa lại nếu nhập muộn). Chỉ sửa nhầm dữ liệu (không phải điều "
        "chuyển) thì xoá dòng để các tháng trước tính theo dữ liệu mới"
    )

    def get_formset(self, request, obj=None, **kwargs):
        # Dòng thêm tay (điều chuyển nhập trước khi có tính năng này) điền sẵn dữ liệu
        # hiện tại — HR chỉ sửa đúng thông tin đã đổi, khỏi để trống thành "không có".
        formset = super().get_formset(request, obj, **kwargs)
        if obj is not None:
            for name, value in (('department', obj.department_id), ('position', obj.position_id),
                                ('doctor_team', obj.doctor_team), ('work_form', obj.work_form)):
                if name in formset.form.base_fields:
                    formset.form.base_fields[name].initial = value
        return formset


@admin.register(EmployeeStatusHistory)
class EmployeeStatusHistoryAdmin(admin.ModelAdmin):
    list_display = ('employee', 'from_status', 'to_status', 'changed_at', 'changed_by')
    list_filter = ('to_status', 'from_status')
    search_fields = ('employee__employee_id', 'employee__full_name')
    raw_id_fields = ('employee', 'changed_by')
    date_hierarchy = 'changed_at'
    ordering = ('-changed_at', '-id')


@admin.register(EmployeeAssignmentHistory)
class EmployeeAssignmentHistoryAdmin(admin.ModelAdmin):
    list_display = ('employee', 'changed_from', 'department', 'position', 'doctor_team', 'work_form', 'created_at')
    list_filter = ('changed_from', 'work_form')
    search_fields = ('employee__employee_id', 'employee__full_name')
    raw_id_fields = ('employee',)
    date_hierarchy = 'changed_from'
    ordering = ('-changed_from', '-id')


class EmployeeAdmin(admin.ModelAdmin):
    form = EmployeeAdminForm
    inlines = [EmployeeAssignmentHistoryInline]
    list_display = ('employee_id', 'full_name', 'position', 'department', 'region', 'subsidiary_legal_entity',
                   'employment_status', 'is_active', 'is_hr', 'start_date', 'subordinate_count', 'probation_rate',
                   'salary_adjustments_preview',)
    list_filter = ('employment_status', 'is_active', 'department', 'position', 'gender', 'is_hr',
                   'region', 'subsidiary_legal_entity', 'work_form', 'education_level', 'marital_status', 'probation_rate',
                   'saturday_afternoon_always_workday', 'overtime_extended_deadline_enabled',)
    search_fields = ('employee_id', 'full_name', 'phone_number', 'personal_email', 
                    'cccd_number', 'social_insurance_number', 'tax_code', 'subsidiary_legal_entity')
    raw_id_fields = ('user', 'position', 'department', 'manager', 'manager_level_2', 'manager_level_3')
    ordering = ('employee_id',)
    
    fieldsets = (
        ('Thông tin tài khoản', {
            'fields': ('user', 'employee_id')
        }),
        ('Thông tin công việc', {
            'fields': ('position', 'department', 'manager', 'manager_level_2', 'manager_level_3',
                      'employment_status', 'is_active', 'is_hr', 'start_date', 'end_date',
                      'termination_reason', 'total_work_months',
                      'saturday_afternoon_always_workday',
                      'overtime_extended_deadline_enabled')
        }),
        ('Thông tin tổ chức', {
            'fields': (
                'region', 'block', 'subsidiary_legal_entity', 'section',
                'sub_department', 'team', 'work_type',
                'rank', 'work_location', 'doctor_team',
            ),
            'classes': ('collapse',)
        }),
        ('Hợp đồng & Học vấn', {
            'fields': ('work_form', 'education_level', 'contract_type', 
                      'probation_end_date', 'probation_months'),
            'classes': ('collapse',)
        }),
        ('Thông tin cá nhân', {
            'fields': ('avatar', 'full_name', 'gender', 'date_of_birth', 
                      'phone_number', 'personal_email', 'facebook_link', 'marital_status')
        }),
        ('Thông tin CCCD/CMND', {
            'fields': ('cccd_number', 'old_id_number', 'cccd_issue_date', 'cccd_issue_place',
                      'birth_place', 'permanent_residence', 'current_address', 'vneid_screenshot'),
            'classes': ('collapse',)
        }),
        ('Thông tin người liên hệ khẩn cấp', {
            'fields': ('emergency_contact_name', 'emergency_contact_relationship', 
                      'emergency_contact_phone', 'emergency_contact_dob',
                      'emergency_contact_occupation', 'emergency_contact_address'),
            'classes': ('collapse',)
        }),
        ('Thông tin ngân hàng', {
            'fields': ('bank_name', 'bank_account'),
            'classes': ('collapse',)
        }),
        ('Thông tin BHXH & Thuế', {
            'fields': ('social_insurance_number', 'tax_code', 'household_code',
                      'insurance_participation', 'insurance_increase_time',
                      'social_insurance_participation_status', 'social_insurance_contribution_level',
                      'social_insurance_start_date', 'social_insurance_end_date'),
            'classes': ('collapse',)
        }),
        ('Lương & Phụ cấp', {
            'fields': ('basic_salary', 'salary_notes', 'allowance', 'allowance_notes',
                      'salary_adjustments', 'contracts_info', 'probation_rate',
                      'other_fixed_monthly_income'),
            'classes': ('collapse',)
        }),
        ('Hồ sơ nhân viên', {
            'fields': ('file_status', 'file_submission_deadline', 'file_submission_date',
                      'file_review_notes',
                      'doc_resume', 'doc_cccd', 'doc_degree', 'doc_health'),
            'classes': ('collapse',)
        }),
        ('Đào tạo', {
            'fields': ('training_presentation_viewed', 'training_presentation_viewed_at'),
            'classes': ('collapse',)
        }),
        ('Ghi chú & Thông tin bổ sung', {
            'fields': ('notes', 'qr_code', 'extra_info', 'reporting_path'),
            'classes': ('collapse',)
        }),
        ('Tài khoản Nurse App', {
            'fields': ('nurse_app_username',),
            'description': 'Tên đăng nhập của nhân viên trên ứng dụng chăm sóc khách hàng (cskh-app).',
        }),
    )
    
    readonly_fields = ('reporting_path',)

    change_list_template = 'admin/employee/employee/change_list.html'

    def get_urls(self):
        urls = super().get_urls()
        custom_urls = [
            path(
                'upload-salary-csv/',
                self.admin_site.admin_view(self.upload_salary_csv_view),
                name='employee_employee_upload_salary_csv',
            ),
            path(
                'upload-legal-entity-csv/',
                self.admin_site.admin_view(self.upload_legal_entity_csv_view),
                name='employee_employee_upload_legal_entity_csv',
            ),
            path(
                'upload-social-insurance/',
                self.admin_site.admin_view(self.upload_social_insurance_view),
                name='employee_employee_upload_social_insurance',
            ),
        ]
        return custom_urls + urls

    def changelist_view(self, request, extra_context=None):
        extra_context = extra_context or {}
        extra_context['upload_salary_csv_url'] = reverse('admin:employee_employee_upload_salary_csv')
        extra_context['upload_legal_entity_csv_url'] = reverse('admin:employee_employee_upload_legal_entity_csv')
        if request.user.is_superuser:
            extra_context['upload_social_insurance_url'] = reverse('admin:employee_employee_upload_social_insurance')
        return super().changelist_view(request, extra_context=extra_context)

    def get_fieldsets(self, request, obj=None):
        fieldsets = super().get_fieldsets(request, obj)
        if request.user.is_superuser:
            return fieldsets
        return tuple(fieldset for fieldset in fieldsets if fieldset[0] != 'Thông tin BHXH & Thuế')

    def upload_social_insurance_view(self, request):
        if not request.user.is_superuser:
            raise PermissionDenied('Bạn không có quyền truy cập chức năng này.')

        form = EmployeeSocialInsuranceUploadForm(request.POST or None, request.FILES or None)
        updated_rows = []
        error_rows = []

        status_choices = {
            key.lower(): key for key, _ in SOCIAL_INSURANCE_PARTICIPATION_STATUS
        }
        status_choices.update({
            label.lower(): key for key, label in SOCIAL_INSURANCE_PARTICIPATION_STATUS
        })

        if request.method == 'POST' and form.is_valid():
            upload_file = form.cleaned_data['file']
            try:
                rows = _read_tabular_rows(upload_file)
                if not rows:
                    raise ValueError('File không có dữ liệu để cập nhật.')

                employee_by_code = {
                    _normalize_employee_code(emp.employee_id): emp
                    for emp in Employee.objects.all().only(
                        'id',
                        'employee_id',
                        'full_name',
                        'social_insurance_number',
                        'tax_code',
                        'household_code',
                        'insurance_participation',
                        'insurance_increase_time',
                        'social_insurance_contribution_level',
                        'social_insurance_participation_status',
                        'social_insurance_start_date',
                        'social_insurance_end_date',
                    )
                }

                for line_number, row in enumerate(rows, start=2):
                    code_raw = _get_cell(row, 'ma_nv', 'employee_id', 'employee_code', 'mã nv', 'mã nhân viên')
                    employee_code = _normalize_employee_code(code_raw)
                    if not employee_code:
                        error_rows.append({
                            'line': line_number,
                            'employee_code': '',
                            'error': 'Thiếu mã nhân viên (ma_nv).',
                        })
                        continue

                    employee = employee_by_code.get(employee_code)
                    if not employee:
                        error_rows.append({
                            'line': line_number,
                            'employee_code': str(code_raw).strip() if code_raw is not None else employee_code,
                            'error': 'Không tìm thấy Employee theo employee_id.',
                        })
                        continue

                    try:
                        status_raw = _normalize_optional_text(_get_cell(
                            row,
                            'social_insurance_participation_status',
                            'trang_thai_tham_gia_bhxh',
                            'trạng thái tham gia bhxh',
                        ))
                        normalized_status = None
                        if status_raw:
                            normalized_status = status_choices.get(status_raw.lower())
                            if not normalized_status:
                                raise ValueError(
                                    'Trạng thái BHXH không hợp lệ. Dùng: PARTICIPATING, NOT_PARTICIPATING, SUSPENDED.'
                                )

                        employee.social_insurance_number = _normalize_optional_text(_get_cell(
                            row,
                            'social_insurance_number',
                            'ma_so_bhxh',
                            'mã số bhxh',
                        ))
                        employee.tax_code = _normalize_optional_text(_get_cell(
                            row,
                            'tax_code',
                            'ma_so_thue_tncn',
                            'mã số thuế tncn',
                        ))
                        employee.household_code = _normalize_optional_text(_get_cell(
                            row,
                            'household_code',
                            'ma_ho_gia_dinh',
                            'mã hộ gia đình',
                        ))
                        employee.insurance_participation = _normalize_optional_text(_get_cell(
                            row,
                            'insurance_participation',
                            'dang_tham_gia_dong_tai_cong_ty',
                            'đang tham gia đóng tại công ty',
                        ))
                        employee.insurance_increase_time = _normalize_optional_text(_get_cell(
                            row,
                            'insurance_increase_time',
                            'thoi_diem_bao_tang',
                            'thời điểm báo tăng',
                        ))
                        employee.social_insurance_contribution_level = _parse_money(
                            _get_cell(
                                row,
                                'social_insurance_contribution_level',
                                'muc_dong_bhxh',
                                'mức đóng bhxh',
                            ),
                            default=None,
                        )
                        employee.social_insurance_participation_status = normalized_status
                        employee.social_insurance_start_date = _parse_date_value(
                            _get_cell(
                                row,
                                'social_insurance_start_date',
                                'ngay_bat_dau_dong_bhxh',
                                'ngày bắt đầu đóng bhxh',
                            ),
                            default=None,
                        )
                        employee.social_insurance_end_date = _parse_date_value(
                            _get_cell(
                                row,
                                'social_insurance_end_date',
                                'ngay_ket_thuc_dong_bhxh',
                                'ngày kết thúc đóng bhxh',
                            ),
                            default=None,
                        )

                        employee.save(update_fields=[
                            'social_insurance_number',
                            'tax_code',
                            'household_code',
                            'insurance_participation',
                            'insurance_increase_time',
                            'social_insurance_contribution_level',
                            'social_insurance_participation_status',
                            'social_insurance_start_date',
                            'social_insurance_end_date',
                            'updated_at',
                        ])

                        updated_rows.append({
                            'line': line_number,
                            'employee_code': employee.employee_id,
                            'employee_name': employee.full_name,
                        })
                    except Exception as exc:
                        error_rows.append({
                            'line': line_number,
                            'employee_code': employee.employee_id,
                            'error': str(exc),
                        })

                self.message_user(
                    request,
                    f'Upload hoàn tất. Thành công: {len(updated_rows)}, lỗi: {len(error_rows)}.',
                    level=messages.SUCCESS if not error_rows else messages.WARNING,
                )
            except Exception as exc:
                self.message_user(request, f'Lỗi xử lý file: {exc}', level=messages.ERROR)

        context = {
            **self.admin_site.each_context(request),
            'opts': self.model._meta,
            'title': 'Upload Excel cập nhật BHXH nhân viên',
            'form': form,
            'updated_rows': updated_rows,
            'error_rows': error_rows,
        }
        return render(request, 'admin/employee/employee/upload_social_insurance.html', context)

    def upload_legal_entity_csv_view(self, request):
        form = EmployeeLegalEntityCSVUploadForm(request.POST or None, request.FILES or None)
        updated_rows = []
        error_rows = []

        if request.method == 'POST' and form.is_valid():
            upload_file = form.cleaned_data['file']
            try:
                raw_bytes = upload_file.read()
                text = raw_bytes.decode('utf-8-sig', errors='ignore')
                sample = text[:2000]
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters='\t,;')
                    delimiter = dialect.delimiter
                except csv.Error:
                    delimiter = '\t' if '\t' in sample else ','

                reader = csv.DictReader(text.splitlines(), delimiter=delimiter)
                if not reader.fieldnames:
                    raise ValueError('Không đọc được header của file CSV.')

                employee_by_code = {
                    _normalize_employee_code(emp.employee_id): emp
                    for emp in Employee.objects.all().only('id', 'employee_id', 'full_name', 'subsidiary_legal_entity')
                }

                for line_number, row in enumerate(reader, start=2):
                    code_raw = _get_cell(row, 'ma_nv', 'employee_id', 'employee_code')
                    employee_code = _normalize_employee_code(code_raw)
                    if not employee_code:
                        error_rows.append({
                            'line': line_number,
                            'employee_code': '',
                            'error': 'Thiếu mã nhân viên (ma_nv).',
                        })
                        continue

                    employee = employee_by_code.get(employee_code)
                    if not employee:
                        error_rows.append({
                            'line': line_number,
                            'employee_code': str(code_raw).strip() if code_raw is not None else employee_code,
                            'error': 'Không tìm thấy Employee theo employee_id.',
                        })
                        continue

                    legal_entity_raw = _get_cell(
                        row,
                        'phap_nhan_con',
                        'phap_nhan',
                        'legal_entity',
                        'subsidiary_legal_entity',
                    )
                    legal_entity = str(legal_entity_raw).strip() if legal_entity_raw is not None else ''
                    if legal_entity == '':
                        legal_entity = None

                    employee.subsidiary_legal_entity = legal_entity
                    employee.save(update_fields=['subsidiary_legal_entity', 'updated_at'])
                    updated_rows.append({
                        'line': line_number,
                        'employee_code': employee.employee_id,
                        'employee_name': employee.full_name,
                        'legal_entity': legal_entity or '',
                    })

                self.message_user(
                    request,
                    f'Upload hoàn tất. Thành công: {len(updated_rows)}, lỗi: {len(error_rows)}.',
                    level=messages.SUCCESS if not error_rows else messages.WARNING,
                )
            except Exception as exc:
                self.message_user(request, f'Lỗi xử lý file: {exc}', level=messages.ERROR)

        context = {
            **self.admin_site.each_context(request),
            'opts': self.model._meta,
            'title': 'Upload CSV cập nhật pháp nhân con',
            'form': form,
            'updated_rows': updated_rows,
            'error_rows': error_rows,
        }
        return render(request, 'admin/employee/employee/upload_legal_entity_csv.html', context)

    def upload_salary_csv_view(self, request):
        form = EmployeeSalaryCSVUploadForm(request.POST or None, request.FILES or None)
        updated_rows = []
        error_rows = []

        if request.method == 'POST' and form.is_valid():
            upload_file = form.cleaned_data['file']
            try:
                raw_bytes = upload_file.read()
                text = raw_bytes.decode('utf-8-sig', errors='ignore')
                sample = text[:2000]
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters='\t,;')
                    delimiter = dialect.delimiter
                except csv.Error:
                    delimiter = '\t' if '\t' in sample else ','

                reader = csv.DictReader(text.splitlines(), delimiter=delimiter)
                if not reader.fieldnames:
                    raise ValueError('Không đọc được header của file CSV.')

                employee_by_code = {
                    _normalize_employee_code(emp.employee_id): emp
                    for emp in Employee.objects.all().only('id', 'employee_id', 'full_name', 'basic_salary', 'salary_adjustments')
                }

                for line_number, row in enumerate(reader, start=2):
                    code_raw = _get_cell(row, 'ma_nv', 'employee_id', 'employee_code')
                    employee_code = _normalize_employee_code(code_raw)
                    if not employee_code:
                        error_rows.append({
                            'line': line_number,
                            'employee_code': '',
                            'error': 'Thiếu mã nhân viên (ma_nv).',
                        })
                        continue

                    employee = employee_by_code.get(employee_code)
                    if not employee:
                        error_rows.append({
                            'line': line_number,
                            'employee_code': str(code_raw).strip() if code_raw is not None else employee_code,
                            'error': 'Không tìm thấy Employee theo employee_id.',
                        })
                        continue

                    try:
                        basic_salary_value = _parse_money(_get_cell(row, 'luong_co_ban'), default=Decimal('0'))
                        insurance_salary = _parse_money(_get_cell(row, 'luong_dong_bao_hiemxh'), default=Decimal('0'))
                        lunch_allowance = _parse_money(_get_cell(row, 'phu_cap_an_trua'), default=Decimal('0'))
                        union_fee = _parse_money(_get_cell(row, 'cong_doan'), default=Decimal('0'))

                        existing_adjustments = dict(employee.salary_adjustments or {})
                        payroll_config = dict(existing_adjustments.get('payroll_config') or {})

                        base_salary_cfg = dict(payroll_config.get('baseSalary') or {})
                        base_salary_cfg['amount'] = float(basic_salary_value)
                        payroll_config['baseSalary'] = base_salary_cfg

                        deductions_cfg = dict(payroll_config.get('deductions') or {})
                        deductions_cfg['unionFee'] = float(union_fee)
                        payroll_config['deductions'] = deductions_cfg

                        lunch_policy = dict(payroll_config.get('lunchAllowancePolicy') or {})
                        lunch_policy['monthlyCap'] = float(lunch_allowance)
                        payroll_config['lunchAllowancePolicy'] = lunch_policy

                        # Save insurance salary directly to social_insurance_contribution_level field
                        update_fields = ['basic_salary', 'salary_adjustments', 'social_insurance_contribution_level', 'updated_at']
                        employee.basic_salary = basic_salary_value
                        employee.social_insurance_contribution_level = insurance_salary
                        employee.salary_adjustments = {
                            'payroll_config': payroll_config,
                        }
                        employee.save(update_fields=update_fields)

                        updated_rows.append({
                            'line': line_number,
                            'employee_code': employee.employee_id,
                            'employee_name': employee.full_name,
                        })
                    except Exception as exc:
                        error_rows.append({
                            'line': line_number,
                            'employee_code': employee.employee_id,
                            'error': str(exc),
                        })

                self.message_user(
                    request,
                    f'Upload hoàn tất. Thành công: {len(updated_rows)}, lỗi: {len(error_rows)}.',
                    level=messages.SUCCESS if not error_rows else messages.WARNING,
                )
            except Exception as exc:
                self.message_user(request, f'Lỗi xử lý file: {exc}', level=messages.ERROR)

        context = {
            **self.admin_site.each_context(request),
            'opts': self.model._meta,
            'title': 'Upload CSV cập nhật lương nhân viên',
            'form': form,
            'updated_rows': updated_rows,
            'error_rows': error_rows,
        }
        return render(request, 'admin/employee/employee/upload_salary_csv.html', context)
    
    def subordinate_count(self, obj):
        return obj.subordinates.count()
    subordinate_count.short_description = "Số cấp dưới"

    def salary_adjustments_preview(self, obj):
        if not obj.salary_adjustments:
            return "—"
        keys = list(obj.salary_adjustments.keys())
        return ", ".join(keys[:3]) + ("…" if len(keys) > 3 else "")
    salary_adjustments_preview.short_description = "Salary adjustments"
    
    def get_queryset(self, request):
        queryset = super().get_queryset(request)
        return queryset.select_related('position', 'department', 'manager')


# -----------------------------
# Inline Admin Classes
# -----------------------------
class EmployeeInline(admin.TabularInline):
    model = Employee
    fields = ('employee_id', 'full_name', 'position', 'subsidiary_legal_entity', 'employment_status')
    extra = 0
    show_change_link = True
    can_delete = False
    
    def has_add_permission(self, request, obj=None):
        return False


class PositionInline(admin.TabularInline):
    model = Position
    fields = ('code', 'title', 'level', 'is_management')
    extra = 0
    show_change_link = True


class DepartmentInline(admin.TabularInline):
    model = Department
    fields = ('code', 'name', 'manager')
    extra = 0
    show_change_link = True
    fk_name = 'parent_department'


# -----------------------------
# Enhanced Admin Classes with Inlines
# -----------------------------
class EnhancedDepartmentAdmin(DepartmentAdmin):
    # Giữ bảng "Đổi lịch làm việc theo ngày hiệu lực" của DepartmentAdmin — khai báo
    # inlines ở đây là THAY HẲN danh sách của lớp cha, thiếu là bảng đó biến mất.
    inlines = [*DepartmentAdmin.inlines, EmployeeInline]


class EnhancedPositionAdmin(PositionAdmin):
    inlines = [EmployeeInline]


# -----------------------------
# Register models
# -----------------------------
admin.site.register(Department, EnhancedDepartmentAdmin)
admin.site.register(Position, EnhancedPositionAdmin)
admin.site.register(Employee, EmployeeAdmin)


# -----------------------------
# BirthdayWish Admin
# -----------------------------
@admin.register(BirthdayWish)
class BirthdayWishAdmin(admin.ModelAdmin):
    list_display = ('id', 'sender', 'recipient', 'year', 'message_preview', 'created_at')
    list_filter = ('year',)
    search_fields = (
        'sender__full_name', 'sender__employee_id',
        'recipient__full_name', 'recipient__employee_id',
        'message',
    )
    raw_id_fields = ('sender', 'recipient')
    ordering = ('-created_at',)
    readonly_fields = ('created_at', 'updated_at')

    fieldsets = (
        ('Thông tin lời chúc', {
            'fields': ('recipient', 'sender', 'year', 'message')
        }),
        ('Thời gian', {
            'fields': ('created_at', 'updated_at'),
            'classes': ('collapse',)
        }),
    )

    def message_preview(self, obj):
        return obj.message[:80] + '…' if len(obj.message) > 80 else obj.message
    message_preview.short_description = "Nội dung"



# -----------------------------
# EmployeeDependent Admin
# -----------------------------
@admin.register(EmployeeDependent)
class EmployeeDependentAdmin(admin.ModelAdmin):
    list_display = ('full_name', 'employee_code', 'employee_name', 'relationship',
                    'date_of_birth', 'dependent_from', 'dependent_to', 'is_active', 'is_deleted')
    list_filter = ('relationship', 'is_active', 'is_deleted')
    actions = ['khoi_phuc', 'xoa_mem']
    search_fields = ('full_name', 'personal_identification_number', 'tax_code',
                     'employee__employee_id', 'employee__full_name')
    raw_id_fields = ('employee',)
    ordering = ('-created_at',)
    readonly_fields = ('created_at', 'updated_at', 'deleted_at', 'deleted_by')

    fieldsets = (
        ('Nhân viên', {
            'fields': ('employee',)
        }),
        ('Thông tin người phụ thuộc', {
            'fields': ('full_name', 'date_of_birth', 'relationship', 'nationality',
                       'personal_identification_number', 'tax_code')
        }),
        ('Thời gian giảm trừ', {
            'fields': ('dependent_from', 'dependent_to', 'is_active', 'notes')
        }),
        ('Xoá mềm', {
            'fields': ('is_deleted', 'deleted_at', 'deleted_by'),
            'description': 'Bản ghi bị xoá ở giao diện chỉ được đánh dấu, không mất dữ liệu. '
                           'Bỏ tick "Đã xoá" để khôi phục.',
        }),
        ('Thời gian', {
            'fields': ('created_at', 'updated_at'),
            'classes': ('collapse',)
        }),
    )

    def get_queryset(self, request):
        return super().get_queryset(request).select_related('employee')

    def employee_code(self, obj):
        return obj.employee.employee_id
    employee_code.short_description = 'Mã NV'

    def employee_name(self, obj):
        return obj.employee.full_name
    employee_name.short_description = 'Họ tên NV'

    @admin.action(description='Khôi phục bản ghi đã xoá')
    def khoi_phuc(self, request, queryset):
        n = queryset.filter(is_deleted=True).update(
            is_deleted=False, deleted_at=None, deleted_by=None)
        self.message_user(request, f'Đã khôi phục {n} người phụ thuộc.')

    @admin.action(description='Xoá mềm')
    def xoa_mem(self, request, queryset):
        from django.utils import timezone as tz
        n = queryset.filter(is_deleted=False).update(
            is_deleted=True, deleted_at=tz.now(), deleted_by=request.user)
        self.message_user(request, f'Đã xoá mềm {n} người phụ thuộc.')



@admin.register(SocialInsurancePeriod)
class SocialInsurancePeriodAdmin(admin.ModelAdmin):
    list_display = ('employee', 'from_month', 'to_month', 'contribution_level', 'employee_rate', 'employer_rate', 'updated_at')
    search_fields = ('employee__employee_id', 'employee__full_name')
    raw_id_fields = ('employee',)
    readonly_fields = ('created_by', 'created_at', 'updated_at')

