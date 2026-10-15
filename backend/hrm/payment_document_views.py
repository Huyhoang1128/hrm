"""API phiếu đề nghị thanh toán / đề xuất mua hàng / đề nghị tạm ứng (chi phí).

/api-hrm/payment-documents/ — chỉ Admin/HCNS. HCNS tự nhập để lưu trữ và rà soát
hằng tháng (không có quy trình duyệt); xem hrm/payment_document_models.py.

Tạo/sửa nhận JSON hoặc multipart: file mới ở field `attachments` (nhiều file),
xoá file cũ qua `remove_attachment_ids`, `items` có thể là chuỗi JSON.
"""
import json
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from io import BytesIO

from django.conf import settings
from django.db import transaction
from django.db.models import Count, Q, Sum
from django.http import HttpResponse
from django.utils import timezone
from rest_framework import serializers, viewsets
from rest_framework.decorators import action
from rest_framework.parsers import FormParser, JSONParser, MultiPartParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from users.utils import get_user_display_name

from .pagination import StandardResultsSetPagination
from .payment_document_models import (
    COMPLETED_LABELS, PAYMENT_DOCUMENT_TYPE, PaymentDocument, PaymentDocumentAttachment,
)
from .permissions import IsAdminOrHR

MAX_ATTACHMENT_SIZE = 20 * 1024 * 1024
ALLOWED_EXTENSIONS = {
    'pdf', 'jpg', 'jpeg', 'png', 'webp', 'heic', 'heif',
    'doc', 'docx', 'xls', 'xlsx', 'csv', 'txt', 'zip', 'rar',
}
TYPE_LABELS = dict(PAYMENT_DOCUMENT_TYPE)


def _attachment_url(obj, request):
    if not obj.file:
        return None
    if getattr(settings, 'CLOUDFLY_ENABLED', False):
        from .utils import get_cloudfly_presigned_url
        return get_cloudfly_presigned_url(str(obj.file))
    return request.build_absolute_uri(obj.file.url) if request else obj.file.url


def _to_decimal(value, field, *, allow_none=False):
    if value in (None, ''):
        if allow_none:
            return None
        raise serializers.ValidationError({field: 'Vui lòng nhập số.'})
    try:
        number = Decimal(str(value).replace(',', '').strip())
    except InvalidOperation:
        raise serializers.ValidationError({field: 'Giá trị phải là số.'})
    if number < 0:
        raise serializers.ValidationError({field: 'Không được âm.'})
    return number


class PaymentDocumentAttachmentSerializer(serializers.ModelSerializer):
    file_url = serializers.SerializerMethodField()

    class Meta:
        model = PaymentDocumentAttachment
        fields = ['id', 'file_name', 'file_url', 'uploaded_at']

    def get_file_url(self, obj):
        return _attachment_url(obj, self.context.get('request'))


class PaymentDocumentSerializer(serializers.ModelSerializer):
    doc_type_display = serializers.CharField(source='get_doc_type_display', read_only=True)
    status_display = serializers.CharField(source='status_label', read_only=True)
    payment_method_display = serializers.CharField(source='get_payment_method_display', read_only=True)
    employee_code = serializers.CharField(source='employee.employee_id', read_only=True, default=None)
    employee_name = serializers.CharField(source='employee.full_name', read_only=True, default=None)
    requester_display = serializers.SerializerMethodField()
    department_name = serializers.CharField(source='department.name', read_only=True, default=None)
    outstanding_amount = serializers.SerializerMethodField()
    reviewed = serializers.SerializerMethodField()
    reviewed_by_name = serializers.SerializerMethodField()
    created_by_name = serializers.SerializerMethodField()
    updated_by_name = serializers.SerializerMethodField()
    attachments = PaymentDocumentAttachmentSerializer(many=True, read_only=True)

    class Meta:
        model = PaymentDocument
        fields = [
            'id', 'code', 'doc_type', 'doc_type_display', 'doc_date',
            'employee', 'employee_code', 'employee_name', 'requester_name', 'requester_display',
            'department', 'department_name',
            'title', 'description', 'amount', 'payee', 'payment_method', 'payment_method_display',
            'bank_account', 'items', 'due_date',
            'status', 'status_display', 'completed_date', 'settled_amount', 'outstanding_amount', 'note',
            'reviewed', 'reviewed_at', 'reviewed_by_name',
            'attachments', 'created_by_name', 'updated_by_name', 'created_at', 'updated_at',
        ]
        read_only_fields = ['code', 'reviewed_at', 'created_at', 'updated_at']
        extra_kwargs = {'amount': {'required': False}}

    # ── hiển thị ──────────────────────────────────────────────────────────
    def get_requester_display(self, obj):
        if obj.employee_id:
            return obj.employee.full_name
        return obj.requester_name

    def get_outstanding_amount(self, obj):
        """Tạm ứng: số còn phải hoàn ứng; loại khác: null."""
        if obj.doc_type != 'EXPENSE_ADVANCE' or obj.status == 'CANCELLED':
            return None
        return str(max((obj.amount or 0) - (obj.settled_amount or 0), 0))

    def get_reviewed(self, obj):
        return obj.reviewed_at is not None

    def get_reviewed_by_name(self, obj):
        return get_user_display_name(obj.reviewed_by) if obj.reviewed_by_id else None

    def get_created_by_name(self, obj):
        return get_user_display_name(obj.created_by) if obj.created_by_id else None

    def get_updated_by_name(self, obj):
        return get_user_display_name(obj.updated_by) if obj.updated_by_id else None

    # ── nhập liệu ─────────────────────────────────────────────────────────
    def to_internal_value(self, data):
        # Multipart gửi mọi thứ dạng chuỗi: '' -> None cho khoá ngoại / ngày / số
        # có thể trống, `items` là chuỗi JSON.
        if hasattr(data, 'getlist'):
            data = {k: data.get(k) for k in data.keys() if k not in ('attachments', 'remove_attachment_ids')}
        else:
            data = dict(data)
        for key in ('employee', 'department', 'due_date', 'completed_date', 'settled_amount'):
            if data.get(key) == '':
                data[key] = None
        if data.get('amount') in ('', None):
            # Mua hàng có danh sách mặt hàng: số tiền do validate() tự tính.
            data.pop('amount', None)
        if isinstance(data.get('items'), str):
            try:
                data['items'] = json.loads(data['items'] or '[]')
            except ValueError:
                raise serializers.ValidationError({'items': 'Danh sách mặt hàng không hợp lệ.'})
        return super().to_internal_value(data)

    def validate_items(self, items):
        if not items:
            return []
        if not isinstance(items, list):
            raise serializers.ValidationError('Danh sách mặt hàng không hợp lệ.')
        cleaned = []
        for idx, item in enumerate(items, 1):
            if not isinstance(item, dict):
                raise serializers.ValidationError(f'Dòng {idx}: dữ liệu không hợp lệ.')
            name = str(item.get('name') or '').strip()
            if not name:
                raise serializers.ValidationError(f'Dòng {idx}: thiếu tên mặt hàng.')
            quantity = _to_decimal(item.get('quantity'), 'items')
            unit_price = _to_decimal(item.get('unit_price'), 'items')
            if quantity <= 0:
                raise serializers.ValidationError(f'Dòng {idx}: số lượng phải lớn hơn 0.')
            cleaned.append({
                'name': name,
                'unit': str(item.get('unit') or '').strip(),
                'quantity': float(quantity),
                'unit_price': float(unit_price),
                'total': float((quantity * unit_price).quantize(Decimal('1'), rounding=ROUND_HALF_UP)),
            })
        return cleaned

    def validate(self, attrs):
        inst = self.instance
        get = lambda key, default=None: attrs[key] if key in attrs else getattr(inst, key, default)  # noqa: E731

        doc_type = get('doc_type')
        employee = get('employee')
        if not employee and not (get('requester_name') or '').strip():
            raise serializers.ValidationError({'employee': 'Chọn người đề nghị hoặc nhập tên người đề nghị.'})
        if employee and not get('department'):
            # Bộ phận để trống thì lấy theo người đề nghị.
            attrs['department'] = employee.department

        items = get('items') or []
        if doc_type != 'PURCHASE_PROPOSAL':
            if 'items' in attrs:
                attrs['items'] = []
        elif items:
            # Có danh sách mặt hàng thì số tiền = tổng thành tiền.
            attrs['amount'] = Decimal(str(sum(i['total'] for i in items)))

        amount = attrs['amount'] if 'amount' in attrs else getattr(inst, 'amount', None)
        if amount is None or Decimal(str(amount)) <= 0:
            raise serializers.ValidationError({'amount': 'Số tiền phải lớn hơn 0.'})

        if doc_type != 'EXPENSE_ADVANCE':
            if 'settled_amount' in attrs:
                attrs['settled_amount'] = None
        elif get('settled_amount') is not None and Decimal(str(get('settled_amount'))) < 0:
            raise serializers.ValidationError({'settled_amount': 'Không được âm.'})

        status = get('status', 'PENDING')
        if status == 'COMPLETED':
            if not get('completed_date'):
                attrs['completed_date'] = timezone.localdate()
        elif 'status' in attrs:
            attrs['completed_date'] = None
        return attrs


def _save_attachments(document, request):
    from rest_framework.exceptions import ValidationError
    files = request.FILES.getlist('attachments')
    for f in files:
        ext = f.name.rsplit('.', 1)[-1].lower() if '.' in f.name else ''
        if ext not in ALLOWED_EXTENSIONS:
            raise ValidationError({'attachments': f'File "{f.name}": định dạng không được hỗ trợ.'})
        if f.size > MAX_ATTACHMENT_SIZE:
            raise ValidationError({'attachments': f'File "{f.name}" vượt quá 20MB.'})
    for f in files:
        if getattr(settings, 'CLOUDFLY_ENABLED', False):
            from .utils import upload_file_to_cloudfly
            key = upload_file_to_cloudfly(f, 'hrm/payment_documents')
            PaymentDocumentAttachment.objects.create(document=document, file=key, file_name=f.name, uploaded_by=request.user)
        else:
            PaymentDocumentAttachment.objects.create(document=document, file=f, file_name=f.name, uploaded_by=request.user)


def _remove_attachments(document, request):
    """remove_attachment_ids: lặp field (multipart), "1,2", "[1,2]" hoặc list JSON."""
    if hasattr(request.data, 'getlist'):
        values = request.data.getlist('remove_attachment_ids')
    else:
        values = request.data.get('remove_attachment_ids') or []
        values = values if isinstance(values, list) else [values]
    ids = []
    for value in values:
        if isinstance(value, str):
            text = value.strip().strip('[]')
            ids.extend(int(x) for x in text.split(',') if x.strip().isdigit())
        elif isinstance(value, int):
            ids.append(value)
    if ids:
        document.attachments.filter(id__in=ids).delete()


class PaymentDocumentViewSet(viewsets.ModelViewSet):
    serializer_class = PaymentDocumentSerializer
    permission_classes = [IsAuthenticated, IsAdminOrHR]
    pagination_class = StandardResultsSetPagination
    parser_classes = [JSONParser, MultiPartParser, FormParser]

    def get_queryset(self):
        qs = (
            PaymentDocument.objects.filter(is_deleted=False)
            .select_related('employee', 'department', 'reviewed_by', 'created_by', 'updated_by')
            .prefetch_related('attachments')
            .order_by('-doc_date', '-id')
        )
        p = self.request.query_params
        if p.get('doc_type'):
            qs = qs.filter(doc_type__in=[t for t in p['doc_type'].split(',') if t])
        if p.get('status'):
            qs = qs.filter(status__in=[s for s in p['status'].split(',') if s])
        month = p.get('month')  # YYYY-MM theo ngày đề nghị
        if month:
            try:
                y, m = (int(x) for x in month.split('-')[:2])
                qs = qs.filter(doc_date__year=y, doc_date__month=m)
            except ValueError:
                qs = qs.none()
        elif p.get('year'):
            qs = qs.filter(doc_date__year=p['year'])
        if p.get('date_from'):
            qs = qs.filter(doc_date__gte=p['date_from'])
        if p.get('date_to'):
            qs = qs.filter(doc_date__lte=p['date_to'])
        if p.get('department'):
            qs = qs.filter(department_id=p['department'])
        if p.get('employee'):
            qs = qs.filter(employee_id=p['employee'])
        if p.get('reviewed') in ('true', 'false'):
            qs = qs.filter(reviewed_at__isnull=(p['reviewed'] == 'false'))
        search = (p.get('search') or '').strip()
        if search:
            qs = qs.filter(
                Q(code__icontains=search) | Q(title__icontains=search) | Q(description__icontains=search)
                | Q(payee__icontains=search) | Q(requester_name__icontains=search)
                | Q(employee__full_name__icontains=search) | Q(employee__employee_id__icontains=search)
            )
        return qs

    @transaction.atomic
    def perform_create(self, serializer):
        document = serializer.save(created_by=self.request.user, updated_by=self.request.user)
        _save_attachments(document, self.request)

    @transaction.atomic
    def perform_update(self, serializer):
        document = serializer.save(updated_by=self.request.user)
        _remove_attachments(document, self.request)
        _save_attachments(document, self.request)

    def update(self, request, *args, **kwargs):
        response = super().update(request, *args, **kwargs)
        # Trả về bản mới nhất (danh sách file đã đổi sau serializer.save()).
        instance = self.get_queryset().get(pk=kwargs['pk'])
        response.data = self.get_serializer(instance).data
        return response

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        self.perform_create(serializer)
        instance = self.get_queryset().get(pk=serializer.instance.pk)
        return Response(self.get_serializer(instance).data, status=201)

    def perform_destroy(self, instance):
        instance.is_deleted = True
        instance.deleted_at = timezone.now()
        instance.deleted_by = self.request.user
        instance.save(update_fields=['is_deleted', 'deleted_at', 'deleted_by', 'updated_at'])

    @action(detail=True, methods=['post'])
    def review(self, request, pk=None):
        """Đánh dấu / bỏ đánh dấu "Đã rà soát" cho 1 phiếu."""
        document = self.get_object()
        reviewed = request.data.get('reviewed', True) in (True, 'true', '1', 1)
        document.reviewed_at = timezone.now() if reviewed else None
        document.reviewed_by = request.user if reviewed else None
        document.save(update_fields=['reviewed_at', 'reviewed_by', 'updated_at'])
        return Response(self.get_serializer(document).data)

    @action(detail=False, methods=['post'], url_path='review-bulk')
    def review_bulk(self, request):
        """Đánh dấu nhiều phiếu (vd cả tháng) đã / chưa rà soát. Body: {ids: [...], reviewed: bool}."""
        ids = request.data.get('ids') or []
        if not isinstance(ids, list) or not ids:
            return Response({'detail': 'Chưa chọn phiếu nào.'}, status=400)
        reviewed = request.data.get('reviewed', True) in (True, 'true', '1', 1)
        qs = self.get_queryset().filter(id__in=ids)
        if reviewed:
            updated = qs.filter(reviewed_at__isnull=True).update(
                reviewed_at=timezone.now(), reviewed_by=request.user, updated_at=timezone.now(),
            )
        else:
            updated = qs.update(reviewed_at=None, reviewed_by=None, updated_at=timezone.now())
        return Response({'updated': updated})

    @action(detail=False, methods=['get'])
    def summary(self, request):
        """Tổng hợp theo bộ lọc hiện tại: số phiếu, tổng tiền (không tính phiếu huỷ), rà soát."""
        qs = self.get_queryset()
        active = qs.exclude(status='CANCELLED')
        by_type = []
        for code, label in PAYMENT_DOCUMENT_TYPE:
            sub = qs.filter(doc_type=code)
            sub_active = sub.exclude(status='CANCELLED')
            row = {
                'doc_type': code,
                'label': label,
                'count': sub.count(),
                'amount': str(sub_active.aggregate(t=Sum('amount'))['t'] or 0),
                'pending': sub.filter(status='PENDING').count(),
                'completed': sub.filter(status='COMPLETED').count(),
                'cancelled': sub.filter(status='CANCELLED').count(),
                'completed_label': COMPLETED_LABELS[code],
            }
            if code == 'EXPENSE_ADVANCE':
                agg = sub_active.aggregate(a=Sum('amount'), s=Sum('settled_amount'))
                row['outstanding'] = str(max((agg['a'] or 0) - (agg['s'] or 0), 0))
            by_type.append(row)
        counts = qs.aggregate(total=Count('id'), reviewed=Count('id', filter=Q(reviewed_at__isnull=False)))
        return Response({
            'total_count': counts['total'],
            'total_amount': str(active.aggregate(t=Sum('amount'))['t'] or 0),
            'reviewed_count': counts['reviewed'],
            'unreviewed_count': counts['total'] - counts['reviewed'],
            'by_type': by_type,
        })

    @action(detail=False, methods=['get'], url_path='export-excel')
    def export_excel(self, request):
        """Xuất danh sách phiếu theo bộ lọc hiện tại ra .xlsx (rà soát hằng tháng)."""
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        qs = self.get_queryset()
        wb = Workbook()
        ws = wb.active
        ws.title = 'Phieu'
        headers = [
            ('Mã phiếu', 16), ('Loại phiếu', 20), ('Ngày đề nghị', 13), ('Người đề nghị', 24),
            ('Mã NV', 11), ('Bộ phận', 20), ('Nội dung', 36), ('Chi tiết mặt hàng', 40),
            ('Người / đơn vị nhận tiền', 26), ('Hình thức TT', 13), ('Số TK / ngân hàng', 24),
            ('Số tiền', 15), ('Hạn', 12), ('Trạng thái', 15), ('Ngày hoàn tất', 13),
            ('Đã hoàn ứng', 14), ('Còn phải hoàn ứng', 16), ('Đã rà soát', 11), ('Người rà soát', 20),
            ('Số file đính kèm', 10), ('Ghi chú', 30),
        ]
        fill = PatternFill(fill_type='solid', fgColor='1E3A5F')
        for col, (title, width) in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=title)
            cell.fill = fill
            cell.font = Font(color='FFFFFF', bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            ws.column_dimensions[cell.column_letter].width = width
        ws.freeze_panes = 'A2'

        money_cols = {12, 16, 17}
        total = Decimal('0')
        row_idx = 1
        for doc in qs:
            row_idx += 1
            items_text = '; '.join(
                f"{i.get('name')} x {i.get('quantity'):g}{(' ' + i['unit']) if i.get('unit') else ''} @ {i.get('unit_price'):,.0f}"
                for i in (doc.items or [])
            )
            outstanding = None
            if doc.doc_type == 'EXPENSE_ADVANCE' and doc.status != 'CANCELLED':
                outstanding = float(max((doc.amount or 0) - (doc.settled_amount or 0), 0))
            if doc.status != 'CANCELLED':
                total += doc.amount or 0
            values = [
                doc.code, TYPE_LABELS.get(doc.doc_type, doc.doc_type), doc.doc_date,
                doc.employee.full_name if doc.employee_id else doc.requester_name,
                doc.employee.employee_id if doc.employee_id else '',
                doc.department.name if doc.department_id else '',
                doc.title + (f"\n{doc.description}" if doc.description else ''), items_text,
                doc.payee, doc.get_payment_method_display() if doc.payment_method else '', doc.bank_account,
                float(doc.amount or 0), doc.due_date, doc.status_label, doc.completed_date,
                float(doc.settled_amount) if doc.settled_amount is not None else None, outstanding,
                'Đã rà soát' if doc.reviewed_at else 'Chưa', get_user_display_name(doc.reviewed_by) if doc.reviewed_by_id else '',
                len(doc.attachments.all()), doc.note,
            ]
            for col, value in enumerate(values, 1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                if isinstance(value, (date, datetime)):
                    cell.number_format = 'DD/MM/YYYY'
                elif col in money_cols and value is not None:
                    cell.number_format = '#,##0'
                if col in (7, 8):
                    cell.alignment = Alignment(wrap_text=True, vertical='top')
        total_row = row_idx + 1
        ws.cell(row=total_row, column=11, value='Tổng (không tính phiếu huỷ)').font = Font(bold=True)
        cell = ws.cell(row=total_row, column=12, value=float(total))
        cell.font = Font(bold=True)
        cell.number_format = '#,##0'

        buffer = BytesIO()
        wb.save(buffer)
        month = request.query_params.get('month') or timezone.localdate().strftime('%Y-%m')
        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = f'attachment; filename="phieu-thanh-toan-{month}.xlsx"'
        return response
