"""Yêu cầu: các cột "Trạng thái vận hành" (2026-09-19), rồi "Phân loại" và
"Tình trạng vật lý" (2026-09-22) trong file Excel export/import tài sản không
còn bị ép chọn từ danh sách dropdown cố định — phải là textfield tự do.

Asset.objects.none() không chạm DB thật (EmptyQuerySet ngắn mạch khi iterate)
(dropdown "Công ty" đọc danh sách đơn vị từ DB nên dùng TestCase).
"""
from django.test import TestCase

from hrm.asset_export import build_template_workbook, build_workbook, COLUMNS
from hrm.models import Asset

FREE_TEXT_COLUMNS = ["Phân loại", "Tình trạng vật lý", "Trạng thái vận hành"]


class DropdownColumnsRemovedTests(TestCase):
    def _validated_columns(self, ws):
        """Trả về set các cột (chữ cái) có áp dụng data validation dropdown."""
        cols = set()
        for dv in ws.data_validations.dataValidation:
            for rng in dv.sqref.ranges:
                cols.add(rng.coord.split(":")[0].rstrip("0123456789"))
        return cols

    def _column_letter(self, header: str) -> str:
        from openpyxl.utils import get_column_letter
        headers = [h for h, _ in COLUMNS]
        return get_column_letter(headers.index(header) + 1)

    def test_free_text_columns_have_no_dropdown_in_template(self):
        wb = build_template_workbook()
        ws = wb["Tài sản"]
        validated = self._validated_columns(ws)
        for header in FREE_TEXT_COLUMNS:
            self.assertNotIn(self._column_letter(header), validated, header)

    def test_free_text_columns_have_no_dropdown_in_real_export(self):
        wb = build_workbook(Asset.objects.none())
        ws = wb["Tài sản"]
        validated = self._validated_columns(ws)
        for header in FREE_TEXT_COLUMNS:
            self.assertNotIn(self._column_letter(header), validated, header)

    def test_other_dropdown_columns_still_validated(self):
        """Đảm bảo việc bỏ dropdown cho 3 cột trên không vô tình xoá luôn
        dropdown của các cột khác không liên quan (VD: "Nhà mạng")."""
        wb = build_template_workbook()
        ws = wb["Tài sản"]
        self.assertIn(self._column_letter("Nhà mạng"), self._validated_columns(ws))


class UserVsHandoverColumnsTests(TestCase):
    """Phát hiện (2026-09-22): "Mã/Tên người bàn giao" trong Excel thực chất
    đang lưu data người SỬ DỤNG (assigned_to), không phải người bàn giao
    (assigned_by). Export giờ có 2 cặp cột riêng biệt, đúng thứ tự."""

    def test_columns_include_both_user_and_handover_pairs_in_order(self):
        headers = [h for h, _ in COLUMNS]
        self.assertIn("Mã người sử dụng", headers)
        self.assertIn("Tên người sử dụng", headers)
        self.assertIn("Mã người bàn giao", headers)
        self.assertIn("Tên người bàn giao", headers)
        # Người sử dụng phải đứng TRƯỚC người bàn giao (thứ tự cột export).
        self.assertLess(headers.index("Mã người sử dụng"), headers.index("Mã người bàn giao"))
        self.assertLess(headers.index("Tên người sử dụng"), headers.index("Tên người bàn giao"))

    def test_template_workbook_builds_without_error_with_new_columns(self):
        wb = build_template_workbook()
        ws = wb["Tài sản"]
        header_row = [c.value for c in ws[1]]
        self.assertIn("Mã người sử dụng", header_row)
        self.assertIn("Mã người bàn giao", header_row)
