"""
"Phân loại" tài sản nhập tự do: gõ gì lưu nấy, KHÔNG
còn quy giá trị ngoài danh sách chuẩn về "Khác".

- loại tự do vẫn có Số lượng + bàn giao nhiều người như "Khác" trước đây
- import lại file cũ: tài sản đã lỡ lưu thành "Khác" + Tên loại chi tiết được
  chuyển sang đúng loại đã gõ, không tạo bản trùng
- lọc / thống kê / xuất Excel / gợi ý nhận cả loại tự do
"""
from io import BytesIO

from django.contrib.auth.models import User
from django.test import TestCase
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from employee.models import Employee
from hrm.asset_import import import_workbook
from hrm.models import Asset

URL = "/api-hrm/assets/"


def _workbook(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Tài sản"
    ws.append(["Mã thiết bị", "Phân loại", "Số lượng", "Mã người sử dụng", "Số lượng bàn giao", "Mô tả / Ghi chú"])
    for r in rows:
        ws.append(r)
    return wb


class FreeTextAssetTypeImportTests(TestCase):
    def setUp(self):
        user = User.objects.create_user(username="nv", password="x")
        self.nv = Employee.objects.create(user=user, employee_id="NV00999", full_name="Nhân Viên")

    def test_loại_tự_do_có_số_lượng_và_bàn_giao_nhiều_đơn_vị(self):
        res = import_workbook(_workbook([
            ["HNHO - CCDCVP - GHE - 018", "CCDCVP", 25, None, None, "Ghế chân quỳ"],
            ["HNHO - CCDCVP - GHE - 018", None, None, "NV00999", 3, None],
        ]))
        self.assertEqual(res["failed"], [])
        asset = Asset.objects.get(name="HNHO - CCDCVP - GHE - 018")
        self.assertEqual(asset.asset_type, "CCDCVP")
        self.assertEqual(asset.get_asset_type_display(), "CCDCVP")
        self.assertEqual((asset.total_quantity, asset.remaining_quantity), (25, 22))
        self.assertTrue(asset.is_multi_holder())
        self.assertNotIn("Phân loại", " ".join(w for d in res["details"] for w in d["warnings"]))

    def test_import_lại_chuyển_tài_sản_đã_lỡ_lưu_thành_khác(self):
        old = Asset.objects.create(
            name="HNHO - CCDCVP - BAN - 002", asset_type="OTHER",
            specifications={"type_name": "CCDCVP", "quantity": 0},
        )
        res = import_workbook(_workbook([["HNHO - CCDCVP - BAN - 002", "CCDCVP", 5, None, None, "Bàn NV"]]))
        self.assertEqual((res["created"], res["updated"]), (0, 1))
        self.assertEqual(Asset.objects.filter(name="HNHO - CCDCVP - BAN - 002").count(), 1)
        old.refresh_from_db()
        self.assertEqual(old.asset_type, "CCDCVP")
        self.assertEqual(old.specifications, {"quantity": 5})
        self.assertEqual(old.total_quantity, 5)

    def test_khác_có_tên_loại_chi_tiết_khác_thì_không_bị_gộp_nhầm(self):
        Asset.objects.create(name="X-1", asset_type="OTHER", specifications={"type_name": "Máy lọc nước"})
        import_workbook(_workbook([["X-1", "CCDCVP", 1, None, None, None]]))
        self.assertEqual(sorted(Asset.objects.filter(name="X-1").values_list("asset_type", flat=True)),
                         ["CCDCVP", "OTHER"])

    def test_loại_chuẩn_giữ_hành_vi_một_người_giữ(self):
        import_workbook(_workbook([["LAP-1", "Laptop", None, "NV00999", 1, None]]))
        asset = Asset.objects.get(name="LAP-1")
        self.assertEqual(asset.asset_type, "LAPTOP")
        self.assertFalse(Asset.is_multi_holder_type(asset.asset_type))


class FreeTextAssetTypeApiTests(TestCase):
    def setUp(self):
        admin = User.objects.create_user(username="admin", password="x", is_staff=True, is_superuser=True)
        self.client = APIClient()
        self.client.force_authenticate(admin)
        Asset.objects.create(name="A-1", asset_type="CCDCVP", specifications={"quantity": 5})
        Asset.objects.create(name="A-2", asset_type="CCDCVP")
        Asset.objects.create(name="A-3", asset_type="Thiết bị điện")
        Asset.objects.create(name="A-4", asset_type="LAPTOP")

    def test_tạo_qua_form_web_lưu_đúng_chữ_đã_gõ(self):
        res = self.client.post(URL, {"name": "B-1", "asset_type": "Thiết bị công nghệ thông tin"}, format="json")
        self.assertEqual(res.status_code, 201, res.data)
        self.assertEqual(res.data["asset_type"], "Thiết bị công nghệ thông tin")
        self.assertEqual(res.data["asset_type_display"], "Thiết bị công nghệ thông tin")

    def test_lọc_danh_sách_theo_loại_tự_do(self):
        res = self.client.get(URL, {"asset_type": "Thiết bị điện"})
        self.assertEqual(res.status_code, 200)
        self.assertEqual([a["name"] for a in res.data["results"]], ["A-3"])

    def test_gợi_ý_phân_loại(self):
        res = self.client.get(f"{URL}type_options/")
        self.assertEqual(res.status_code, 200)
        opts = {o["value"]: o for o in res.data}
        self.assertEqual(opts["CCDCVP"]["count"], 2)
        self.assertEqual(opts["Thiết bị điện"]["label"], "Thiết bị điện")
        self.assertEqual(opts["LAPTOP"]["label"], "Laptop")
        self.assertEqual(opts["SIM"]["count"], 0)  # loại chuẩn chưa dùng vẫn gợi ý
        self.assertEqual(res.data[0]["value"], "CCDCVP")  # dùng nhiều nhất lên đầu

    def test_thống_kê_gồm_loại_tự_do(self):
        res = self.client.get(f"{URL}stats/")
        self.assertEqual(res.status_code, 200)
        stats = {t["asset_type"]: t for t in res.data["type_stats"]}
        self.assertEqual(stats["CCDCVP"]["count"], 2)
        self.assertEqual(stats["Thiết bị điện"]["asset_type_display"], "Thiết bị điện")

    def test_xuất_excel_lọc_nhiều_loại(self):
        res = self.client.get(f"{URL}export_excel/", {"asset_type": "CCDCVP,laptop"})
        self.assertEqual(res.status_code, 200)
        ws = load_workbook(BytesIO(res.content))["Tài sản"]
        names = sorted(r[0] for r in ws.iter_rows(min_row=2, values_only=True) if r[0])
        self.assertEqual(names, ["A-1", "A-2", "A-4"])
        types = {r[0]: r[1] for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
        self.assertEqual(types["A-1"], "CCDCVP")
        self.assertEqual(types["A-4"], "Laptop")
