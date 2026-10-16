"""Placeholder hợp đồng: số tiền / số tháng bằng chữ, đồng bộ sau override,
và thay placeholder ở header/footer riêng của trang đầu."""
from django.test import SimpleTestCase
from docx import Document

from hrm.views import (
    _amount_to_vietnamese_words, _month_count_to_vietnamese_words, _parse_vnd_amount,
    _sync_word_placeholders, replace_all_placeholders,
)


class VietnameseWordsTests(SimpleTestCase):
    def test_amounts(self):
        self.assertEqual(_amount_to_vietnamese_words(30_000_000), 'Ba mươi triệu đồng')
        self.assertEqual(_amount_to_vietnamese_words(6_500_000), 'Sáu triệu năm trăm nghìn đồng')
        self.assertEqual(_amount_to_vietnamese_words(1_005_000), 'Một triệu không trăm lẻ năm nghìn đồng')
        self.assertEqual(_amount_to_vietnamese_words(21_000), 'Hai mươi mốt nghìn đồng')
        self.assertEqual(_amount_to_vietnamese_words(0), '')
        self.assertEqual(_amount_to_vietnamese_words('abc'), '')

    def test_months(self):
        self.assertEqual(_month_count_to_vietnamese_words(2), 'hai')
        self.assertEqual(_month_count_to_vietnamese_words('02'), 'hai')
        self.assertEqual(_month_count_to_vietnamese_words(12), 'mười hai')
        self.assertEqual(_month_count_to_vietnamese_words(None), '')

    def test_parse_vnd(self):
        self.assertEqual(_parse_vnd_amount('6.000.000'), 6_000_000)
        self.assertEqual(_parse_vnd_amount('6,000,000'), 6_000_000)
        self.assertEqual(_parse_vnd_amount(''), 0)

    def test_words_follow_overridden_numbers(self):
        p = {'{{luong_co_ban}}': '8.000.000', '{{luong_co_ban_bang_chu}}': 'sai',
             '{{so_thang_thu_viec}}': '3', '{{so_thang_thu_viec_bang_chu}}': 'hai'}
        _sync_word_placeholders(p)
        self.assertEqual(p['{{luong_co_ban_bang_chu}}'], 'Tám triệu đồng')
        self.assertEqual(p['{{so_thang_thu_viec_bang_chu}}'], 'ba')


class FirstPageFooterTests(SimpleTestCase):
    def test_first_page_footer_is_filled(self):
        doc = Document()
        doc.add_paragraph('Họ tên: {{ho_ten}}')
        section = doc.sections[0]
        section.different_first_page_header_footer = True
        section.first_page_footer.paragraphs[0].text = 'Mã NV {{mnv}}'
        section.footer.paragraphs[0].text = 'Trang sau {{mnv}}'
        replace_all_placeholders(doc, {'{{ho_ten}}': 'Nguyễn Văn A', '{{mnv}}': 'NV00001'})
        self.assertEqual(doc.paragraphs[0].text, 'Họ tên: Nguyễn Văn A')
        self.assertEqual(section.first_page_footer.paragraphs[0].text, 'Mã NV NV00001')
        self.assertEqual(section.footer.paragraphs[0].text, 'Trang sau NV00001')
