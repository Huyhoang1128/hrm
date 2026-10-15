"""
hrm/test_payment_documents.py
═════════════════════════════
Phiếu đề nghị thanh toán / đề xuất mua hàng / đề nghị tạm ứng — HCNS tự nhập
để lưu trữ và rà soát hằng tháng, không có quy trình duyệt.
"""
import datetime as dt
import json
import shutil
import tempfile
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from openpyxl import load_workbook
from rest_framework.test import APIClient

from employee.models import Department, Employee
from hrm.models import PaymentDocument

URL = '/api-hrm/payment-documents/'
MEDIA = tempfile.mkdtemp(prefix='paydoc-test-')


def _employee(username, code, **kw):
    user = User.objects.create_user(username=username, password='x', is_staff=kw.pop('is_staff', False))
    return user, Employee.objects.create(user=user, employee_id=code, full_name=kw.pop('full_name', code), **kw)


@override_settings(CLOUDFLY_ENABLED=False, MEDIA_ROOT=MEDIA)
class PaymentDocumentApiTests(TestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(MEDIA, ignore_errors=True)

    def setUp(self):
        self.dept = Department.objects.create(name='Phòng khám 1', code='PK1')
        self.hr_user, self.hr = _employee('hr', 'NV90001', full_name='Lê HCNS', is_hr=True)
        self.staff_user, self.staff = _employee('nv', 'NV90002', full_name='Trần Nhân Viên', department=self.dept)
        self.admin_user = User.objects.create_user(username='admin', password='x', is_staff=True)
        self.client = APIClient()
        self.client.force_authenticate(self.hr_user)

    def _create(self, **data):
        payload = {
            'doc_type': 'PAYMENT_REQUEST', 'doc_date': '2026-10-05', 'employee': self.staff.id,
            'title': 'Thanh toán tiền điện tháng 9', 'amount': '1250000', 'payee': 'EVN Hà Nội',
            'payment_method': 'TRANSFER',
        }
        payload.update(data)
        return self.client.post(URL, payload, format='json')

    # ── quyền ────────────────────────────────────────────────────────────
    def test_only_admin_or_hr(self):
        self.client.force_authenticate(self.staff_user)
        self.assertEqual(self.client.get(URL).status_code, 403)
        self.assertEqual(self._create().status_code, 403)
        self.client.force_authenticate(self.admin_user)
        self.assertEqual(self.client.get(URL).status_code, 200)

    # ── tạo ──────────────────────────────────────────────────────────────
    def test_create_payment_request_generates_code_and_department(self):
        res = self._create()
        self.assertEqual(res.status_code, 201, res.data)
        self.assertEqual(res.data['code'], 'DNTT-202610-0001')
        self.assertEqual(res.data['department_name'], 'Phòng khám 1')
        self.assertEqual(res.data['status_display'], 'Chờ xử lý')
        self.assertEqual(res.data['created_by_name'], 'Lê HCNS')
        second = self._create(doc_date='2026-10-20')
        self.assertEqual(second.data['code'], 'DNTT-202610-0002')
        other_month = self._create(doc_date='2026-11-01')
        self.assertEqual(other_month.data['code'], 'DNTT-202611-0001')
        # Form gửi bộ phận rỗng vẫn lấy theo người đề nghị.
        blank_dept = self._create(department=None)
        self.assertEqual(blank_dept.data['department_name'], 'Phòng khám 1')

    def test_requester_required(self):
        res = self._create(employee=None)
        self.assertEqual(res.status_code, 400)
        res = self._create(employee=None, requester_name='Đối tác ngoài')
        self.assertEqual(res.status_code, 201, res.data)
        self.assertEqual(res.data['requester_display'], 'Đối tác ngoài')

    def test_amount_must_be_positive(self):
        self.assertEqual(self._create(amount='0').status_code, 400)

    def test_purchase_items_set_amount(self):
        res = self._create(
            doc_type='PURCHASE_PROPOSAL', amount='1',
            items=[{'name': 'Găng tay', 'unit': 'hộp', 'quantity': 10, 'unit_price': 85000},
                   {'name': 'Khẩu trang', 'quantity': '2', 'unit_price': '45000'}],
        )
        self.assertEqual(res.status_code, 201, res.data)
        self.assertEqual(res.data['code'], 'DXMH-202610-0001')
        self.assertEqual(Decimal(res.data['amount']), Decimal('940000'))
        self.assertEqual(res.data['items'][0]['total'], 850000.0)
        bad = self._create(doc_type='PURCHASE_PROPOSAL', items=[{'name': '', 'quantity': 1, 'unit_price': 1}])
        self.assertEqual(bad.status_code, 400)

    def test_items_dropped_for_other_types(self):
        res = self._create(items=[{'name': 'X', 'quantity': 1, 'unit_price': 1}])
        self.assertEqual(res.data['items'], [])

    def test_expense_advance_outstanding_and_completion(self):
        res = self._create(doc_type='EXPENSE_ADVANCE', title='Tạm ứng công tác', amount='5000000', settled_amount='3200000')
        self.assertEqual(res.status_code, 201, res.data)
        self.assertEqual(Decimal(res.data['outstanding_amount']), Decimal('1800000'))
        doc_id = res.data['id']
        res = self.client.patch(f'{URL}{doc_id}/', {'status': 'COMPLETED', 'settled_amount': '5000000'}, format='json')
        self.assertEqual(res.data['status_display'], 'Đã hoàn ứng')
        self.assertIsNotNone(res.data['completed_date'])
        self.assertEqual(Decimal(res.data['outstanding_amount']), Decimal('0'))
        res = self.client.patch(f'{URL}{doc_id}/', {'status': 'PENDING'}, format='json')
        self.assertIsNone(res.data['completed_date'])

    # ── file đính kèm (multipart) ────────────────────────────────────────
    def test_multipart_attachments_add_and_remove(self):
        res = self.client.post(URL, {
            'doc_type': 'PURCHASE_PROPOSAL', 'doc_date': '2026-10-05', 'employee': str(self.staff.id),
            'department': '', 'due_date': '', 'title': 'Mua vật tư', 'amount': '',
            'items': json.dumps([{'name': 'Bông', 'quantity': 3, 'unit_price': 20000}]),
            'attachments': [
                SimpleUploadedFile('bao-gia.pdf', b'%PDF-1.4 test', content_type='application/pdf'),
                SimpleUploadedFile('hoa-don.jpg', b'jpgdata', content_type='image/jpeg'),
            ],
        }, format='multipart')
        self.assertEqual(res.status_code, 201, res.data)
        self.assertEqual(Decimal(res.data['amount']), Decimal('60000'))
        self.assertEqual([a['file_name'] for a in res.data['attachments']], ['bao-gia.pdf', 'hoa-don.jpg'])
        self.assertTrue(res.data['attachments'][0]['file_url'])
        remove_id = res.data['attachments'][0]['id']
        res = self.client.patch(f"{URL}{res.data['id']}/", {
            'remove_attachment_ids': str(remove_id),
            'attachments': [SimpleUploadedFile('ky-nhan.png', b'png', content_type='image/png')],
        }, format='multipart')
        self.assertEqual(res.status_code, 200, res.data)
        self.assertEqual([a['file_name'] for a in res.data['attachments']], ['hoa-don.jpg', 'ky-nhan.png'])

    def test_rejects_unsupported_file(self):
        res = self.client.post(URL, {
            'doc_type': 'PAYMENT_REQUEST', 'doc_date': '2026-10-05', 'employee': str(self.staff.id),
            'title': 'X', 'amount': '1000',
            'attachments': [SimpleUploadedFile('virus.exe', b'MZ', content_type='application/octet-stream')],
        }, format='multipart')
        self.assertEqual(res.status_code, 400)
        self.assertFalse(PaymentDocument.objects.exists())  # tạo phiếu được rollback

    # ── lọc / rà soát / tổng hợp / xuất ──────────────────────────────────
    def _seed_month(self):
        a = self._create().data
        b = self._create(doc_type='PURCHASE_PROPOSAL', amount='2000000', title='Mua máy in').data
        c = self._create(doc_type='EXPENSE_ADVANCE', amount='3000000', settled_amount='1000000', title='Tạm ứng').data
        d = self._create(amount='999', status='CANCELLED', title='Huỷ').data
        self._create(doc_date='2026-09-28', title='Tháng trước')
        return a, b, c, d

    def test_month_and_type_filters(self):
        self._seed_month()
        self.assertEqual(self.client.get(URL, {'month': '2026-10'}).data['count'], 4)
        self.assertEqual(self.client.get(URL, {'month': '2026-09'}).data['count'], 1)
        res = self.client.get(URL, {'month': '2026-10', 'doc_type': 'PURCHASE_PROPOSAL,EXPENSE_ADVANCE'})
        self.assertEqual(res.data['count'], 2)
        self.assertEqual(self.client.get(URL, {'search': 'máy in'}).data['count'], 1)
        self.assertEqual(self.client.get(URL, {'search': 'NV90002', 'month': '2026-10'}).data['count'], 4)

    def test_review_single_and_bulk(self):
        a, b, c, d = self._seed_month()
        res = self.client.post(f"{URL}{a['id']}/review/", {'reviewed': True}, format='json')
        self.assertTrue(res.data['reviewed'])
        self.assertEqual(res.data['reviewed_by_name'], 'Lê HCNS')
        self.assertEqual(self.client.get(URL, {'month': '2026-10', 'reviewed': 'false'}).data['count'], 3)
        res = self.client.post(f'{URL}review-bulk/', {'ids': [b['id'], c['id'], d['id']], 'reviewed': True}, format='json')
        self.assertEqual(res.data['updated'], 3)
        self.assertEqual(self.client.get(URL, {'month': '2026-10', 'reviewed': 'false'}).data['count'], 0)
        self.client.post(f'{URL}review-bulk/', {'ids': [a['id']], 'reviewed': False}, format='json')
        self.assertEqual(self.client.get(URL, {'month': '2026-10', 'reviewed': 'false'}).data['count'], 1)

    def test_summary_excludes_cancelled(self):
        self._seed_month()
        data = self.client.get(f'{URL}summary/', {'month': '2026-10'}).data
        self.assertEqual(data['total_count'], 4)
        self.assertEqual(Decimal(data['total_amount']), Decimal('6250000'))
        by_type = {r['doc_type']: r for r in data['by_type']}
        self.assertEqual(by_type['PAYMENT_REQUEST']['count'], 2)
        self.assertEqual(by_type['PAYMENT_REQUEST']['cancelled'], 1)
        self.assertEqual(Decimal(by_type['PAYMENT_REQUEST']['amount']), Decimal('1250000'))
        self.assertEqual(Decimal(by_type['EXPENSE_ADVANCE']['outstanding']), Decimal('2000000'))
        self.assertEqual(data['unreviewed_count'], 4)

    def test_soft_delete(self):
        doc = self._create().data
        self.assertEqual(self.client.delete(f"{URL}{doc['id']}/").status_code, 204)
        self.assertEqual(self.client.get(URL).data['count'], 0)
        obj = PaymentDocument.objects.get(pk=doc['id'])
        self.assertTrue(obj.is_deleted)
        self.assertEqual(obj.deleted_by, self.hr_user)

    def test_export_excel(self):
        self._seed_month()
        res = self.client.get(f'{URL}export-excel/', {'month': '2026-10'})
        self.assertEqual(res.status_code, 200)
        ws = load_workbook(BytesIO(res.content)).active
        rows = list(ws.iter_rows(values_only=True))
        self.assertEqual(rows[0][0], 'Mã phiếu')
        self.assertEqual(len(rows), 1 + 4 + 1)  # header + 4 phiếu + dòng tổng
        self.assertEqual(rows[-1][11], 6250000)
        purchase = next(r for r in rows if r[1] == 'Đề xuất mua hàng')
        self.assertEqual(purchase[3], 'Trần Nhân Viên')
