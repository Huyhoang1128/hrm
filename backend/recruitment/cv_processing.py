"""
Utilities for uploading, extracting and parsing candidate CV files.
"""
from __future__ import annotations

import json
import logging
import os
import re
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Tuple

from django.conf import settings

from hrm.utils import get_cloudfly_presigned_url, upload_file_to_cloudfly


logger = logging.getLogger(__name__)


def _clean_text(value: str) -> str:
    return re.sub(r'\s+', ' ', value or '').strip()


def _normalize_gender(value: Any) -> str:
    if not value:
        return ''

    normalized = str(value).strip().lower()
    if normalized in {'m', 'male', 'nam', 'man', 'nam giới'}:
        return 'M'
    if normalized in {'f', 'female', 'nu', 'nữ', 'nữ giới'}:
        return 'F'
    if normalized in {'o', 'other', 'khac', 'khác'}:
        return 'O'
    return ''


def _normalize_date(value: Any) -> str:
    if not value:
        return ''

    if isinstance(value, str):
        text = value.strip()
        for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%Y/%m/%d'):
            try:
                return datetime.strptime(text, fmt).date().isoformat()
            except ValueError:
                continue
    return ''


def _normalize_int(value: Any) -> int | None:
    if value in (None, '', 'null'):
        return None
    try:
        return int(float(str(value).replace(',', '.')))
    except (TypeError, ValueError):
        return None


def _extract_text_from_pdf(file_obj) -> str:
    from pypdf import PdfReader

    file_obj.seek(0)
    reader = PdfReader(file_obj)
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or '')
    return '\n'.join(pages).strip()


def _extract_text_from_docx(file_obj) -> str:
    from docx import Document

    file_obj.seek(0)
    document = Document(file_obj)
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]

    for table in document.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                paragraphs.append(' | '.join(row_text))

    return '\n'.join(paragraphs).strip()


def _extract_text_from_doc(file_obj) -> str:
    """
    Best-effort legacy .doc extraction.

    The backend environment may not have antiword/soffice installed, so we try
    a few safe fallbacks before giving up.
    """
    file_obj.seek(0)
    with tempfile.NamedTemporaryFile(suffix='.doc', delete=False) as tmp_file:
        tmp_file.write(file_obj.read())
        tmp_path = Path(tmp_file.name)

    try:
        for command in (
            ['antiword', str(tmp_path)],
            ['soffice', '--headless', '--convert-to', 'txt', '--outdir', str(tmp_path.parent), str(tmp_path)],
        ):
            try:
                import subprocess

                result = subprocess.run(command, capture_output=True, text=True, check=False)
                if result.returncode == 0 and result.stdout.strip():
                    return result.stdout.strip()
            except FileNotFoundError:
                continue
            except Exception:
                continue
    finally:
        try:
            tmp_path.unlink(missing_ok=True)
        except Exception:
            pass

    return ''


def extract_text_from_cv(file_obj) -> Tuple[str, List[str]]:
    """Extract text from PDF/DOCX/DOC files."""
    warnings: List[str] = []
    filename = (getattr(file_obj, 'name', '') or '').lower()

    try:
        if filename.endswith('.pdf'):
            return _extract_text_from_pdf(file_obj), warnings
        if filename.endswith('.docx'):
            return _extract_text_from_docx(file_obj), warnings
        if filename.endswith('.doc'):
            text = _extract_text_from_doc(file_obj)
            if not text:
                warnings.append('Không thể trích xuất văn bản từ file .doc; nên chuyển sang PDF hoặc DOCX.')
            return text, warnings
    except Exception as exc:
        logger.exception('Failed to extract CV text')
        warnings.append(f'Không thể trích xuất text từ CV: {exc}')
        return '', warnings

    warnings.append('Định dạng CV chưa được hỗ trợ đầy đủ.')
    return '', warnings


def _heuristic_candidate_extract(text: str) -> Dict[str, Any]:
    email_match = re.search(r'[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}', text, flags=re.IGNORECASE)
    phone_match = re.search(r'(?:\+?84|0)\s*(?:\d[\s.-]*){8,10}', text)
    lines = [line.strip('•*-\t ').strip() for line in text.splitlines() if line.strip()]
    email_value = email_match.group(0).lower() if email_match else ''

    full_name = ''
    for line in lines[:8]:
        lower_line = line.lower()
        if len(line) <= 60 and email_value not in lower_line and sum(ch.isalpha() for ch in line) >= 4:
            full_name = line
            break

    return {
        'full_name': full_name,
        'email': email_match.group(0) if email_match else '',
        'phone': phone_match.group(0).replace(' ', '').replace('.', '').replace('-', '') if phone_match else '',
        'gender': '',
        'date_of_birth': '',
        'address': '',
        'skills': '',
        'experience_years': None,
        'experience_summary': '',
        'overall_evaluation': '',
        'education': '',
        'current_company': '',
        'current_position': '',
        'expected_salary': None,
        'notes': '',
        'tags': '',
        'source': 'HR_PERSONAL',
    }


def _merge_candidate_payload(base: Dict[str, Any], parsed: Dict[str, Any]) -> Dict[str, Any]:
    merged = dict(base)
    for key, value in parsed.items():
        if value in (None, '', [], {}):
            continue
        merged[key] = value
    merged['gender'] = _normalize_gender(merged.get('gender'))
    merged['date_of_birth'] = _normalize_date(merged.get('date_of_birth'))
    merged['experience_years'] = _normalize_int(merged.get('experience_years'))
    merged['expected_salary'] = _normalize_int(merged.get('expected_salary'))
    merged['source'] = 'HR_PERSONAL'
    return merged


def _parse_with_openai(text: str) -> Tuple[Dict[str, Any], List[str]]:
    api_key = getattr(settings, 'OPENAI_API_KEY', '') or os.environ.get('OPENAI_API_KEY', '')
    model = getattr(settings, 'OPENAI_MODEL', 'gpt-4o-mini') or 'gpt-4o-mini'
    if not api_key:
        return {}, ['OPENAI_API_KEY chưa được cấu hình, chỉ dùng kết quả bóc tách cơ bản.']

    try:
        from openai import OpenAI
    except Exception as exc:
        logger.exception('OpenAI package is not available')
        return {}, [f'Không thể sử dụng OpenAI SDK: {exc}']

    client = OpenAI(api_key=api_key)
    truncated_text = text[:24000]

    system_prompt = (
        'Bạn là trợ lý tuyển dụng. Hãy đọc nội dung CV và trả về JSON hợp lệ, '
        'không kèm markdown, chỉ gồm các khóa: full_name, email, phone, gender, '
        'date_of_birth, address, skills, experience_years, experience_summary, '
        'overall_evaluation, education, current_company, current_position, '
        'expected_salary, notes, tags. '
        'Các giá trị thiếu để rỗng hoặc null. gender dùng M/F/O nếu có thể suy ra. '
        'date_of_birth trả về định dạng YYYY-MM-DD nếu xác định được. '
        'experience_years và expected_salary là số. overall_evaluation là đánh giá '
        'tổng quan ngắn gọn bằng tiếng Việt về mức độ phù hợp của ứng viên, điểm mạnh, '
        'điểm cần làm rõ và khuyến nghị bước tiếp theo dựa trên CV.'
    )

    user_prompt = f'CV text:\n{truncated_text}'

    response = client.chat.completions.create(
        model=model,
        temperature=0.1,
        response_format={'type': 'json_object'},
        messages=[
            {'role': 'system', 'content': system_prompt},
            {'role': 'user', 'content': user_prompt},
        ],
    )

    content = response.choices[0].message.content or '{}'
    try:
        return json.loads(content), []
    except json.JSONDecodeError:
        return {}, ['OpenAI trả về JSON không hợp lệ.']


def process_candidate_cv(file_obj) -> Dict[str, Any]:
    """Upload CV to S3, extract text and parse candidate data."""
    uploaded_key = upload_file_to_cloudfly(file_obj, 'recruitment/cv')
    cv_url = get_cloudfly_presigned_url(uploaded_key)

    file_obj.seek(0)
    extracted_text, warnings = extract_text_from_cv(file_obj)

    heuristic_data = _heuristic_candidate_extract(extracted_text) if extracted_text else {'source': 'UPLOAD'}
    openai_data, openai_warnings = _parse_with_openai(extracted_text) if extracted_text else ({}, [])
    parsed_data = _merge_candidate_payload(heuristic_data, openai_data)

    return {
        'cv_file': uploaded_key,
        'cv_file_url': cv_url,
        'extracted_text': extracted_text,
        'candidate': parsed_data,
        'warnings': warnings + openai_warnings,
        'source': 'HR_PERSONAL',
    }
